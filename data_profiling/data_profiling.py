from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_FILE = PROJECT_ROOT / "data" / "raw" / "PBJ_Daily_Nurse_Staffing_Q2_2024.csv"
REPORT_DIR = PROJECT_ROOT / "data_profiling"
REPORT_DIR.mkdir(exist_ok=True)
REPORT_FILE = REPORT_DIR / "data_profiling_report.xlsx"

# Main folder locations used by the script.


def load_data():
    """Load the PBJ staffing dataset."""
    # Read the source CSV file into a DataFrame.
    return pd.read_csv(DATA_FILE, encoding="cp1252")


def create_overview(df):
    """Create a simple overview of the dataset."""
    overview = pd.DataFrame(
        {
            "Metric": [
                "Dataset Name",
                "Total Rows",
                "Total Columns",
                "Numeric Columns",
                "Categorical Columns",
                "Duplicate Rows",
                "Total Missing Cells",
                "Memory Usage (MB)",
            ],
            "Value": [
                DATA_FILE.name,
                len(df),
                len(df.columns),
                len(df.select_dtypes(include="number").columns),
                len(df.select_dtypes(exclude="number").columns),
                df.duplicated().sum(),
                df.isnull().sum().sum(),
                round(df.memory_usage(deep=True).sum() / 1024**2, 2),
            ],
        }
    )
    return overview


def create_data_dictionary(df):
    """Create a data dictionary for the dataset."""
    return pd.DataFrame(
        {
            "Column Name": df.columns,
            "Data Type": df.dtypes.astype(str),
            "Missing Count": df.isnull().sum().values,
            "Missing %": (df.isnull().mean() * 100).round(2).values,
            "Unique Values": df.nunique().values,
        }
    )


def create_missing_values_summary(df):
    """Show which columns have the most missing values."""
    missing = pd.DataFrame(
        {
            "Column": df.columns,
            "Missing Count": df.isnull().sum().values,
            "Missing %": (df.isnull().mean() * 100).round(2).values,
        }
    )
    return missing.sort_values(by="Missing %", ascending=False)


def create_summary_statistics(df):
    """Create summary statistics for the dataset."""
    return df.describe(include="all").transpose()


def create_data_quality_report(df):
    """Create a simple data quality report."""
    checks = [{"Check": "Duplicate Rows", "Result": df.duplicated().sum()}]

    for column in df.columns:
        checks.append(
            {
                "Check": f"Missing Values - {column}",
                "Result": df[column].isnull().sum(),
            }
        )

    return pd.DataFrame(checks)


def create_sample_data(df):
    """Return the first 20 rows as a sample."""
    return df.head(20)


def format_worksheet(worksheet):
    """Apply simple formatting to an Excel worksheet."""
    header_fill = PatternFill(fill_type="solid", fgColor="4F81BD")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.alignment = body_alignment

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    worksheet.row_dimensions[1].height = 30
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is not None:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length

        adjusted_width = max_length + 4
        adjusted_width = min(max(adjusted_width, 12), 45)
        worksheet.column_dimensions[column_letter].width = adjusted_width

    headers = [cell.value for cell in worksheet[1]]

    for col_index, header in enumerate(headers, start=1):
        if header and "%" in str(header):
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row, column=col_index).number_format = "0.00"

    if "Missing %" in headers:
        column_index = headers.index("Missing %") + 1
        column_letter = get_column_letter(column_index)
        data_range = f"{column_letter}2:{column_letter}{worksheet.max_row}"

        worksheet.conditional_formatting.add(
            data_range,
            CellIsRule(operator="equal", formula=["0"], fill=PatternFill(fill_type="solid", fgColor="C6EFCE")),
        )
        worksheet.conditional_formatting.add(
            data_range,
            CellIsRule(operator="between", formula=["0.0001", "0.10"], fill=PatternFill(fill_type="solid", fgColor="FFEB9C")),
        )
        worksheet.conditional_formatting.add(
            data_range,
            CellIsRule(operator="greaterThan", formula=["0.10"], fill=PatternFill(fill_type="solid", fgColor="FFC7CE")),
        )


def create_excel_report(df):
    """Create the Excel profiling report."""
    # Build each report section before writing the workbook.
    overview = create_overview(df)
    data_dictionary = create_data_dictionary(df)
    missing_values = create_missing_values_summary(df)
    summary_stats = create_summary_statistics(df)
    quality_report = create_data_quality_report(df)
    sample_data = create_sample_data(df)

    with pd.ExcelWriter(REPORT_FILE, engine="openpyxl") as writer:
        # Write each section to its own worksheet.
        overview.to_excel(writer, sheet_name="Overview", index=False)
        data_dictionary.to_excel(writer, sheet_name="Data Dictionary", index=False)
        missing_values.to_excel(writer, sheet_name="Missing Values", index=False)
        summary_stats.to_excel(writer, sheet_name="Summary Statistics")
        quality_report.to_excel(writer, sheet_name="Data Quality", index=False)
        sample_data.to_excel(writer, sheet_name="Sample Data", index=False)

        for worksheet in writer.book.worksheets:
            format_worksheet(worksheet)


def main():
    print("Loading dataset...")
    df = load_data()
    print("Generating profiling report...")
    create_excel_report(df)
    print("\nDone!")
    print(f"\nReport saved to:\n{REPORT_FILE}")


if __name__ == "__main__":
    main()

